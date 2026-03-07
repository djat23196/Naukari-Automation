from __future__ import annotations

import csv
import os
from datetime import datetime
from typing import Set

from playwright.sync_api import Page, TimeoutError as PlaywrightTimeout

from openpyxl import load_workbook

from chatbot import handle_chatbot
from login import human_delay
from search import JobListing

CSV_FILE = "applied_jobs.csv"
CSV_HEADERS = ["job_id", "title", "company", "link", "status", "timestamp"]
EXTERNAL_CSV = "external_apply_jobs.csv"
EXTERNAL_HEADERS = ["title", "company", "link", "timestamp"]
EXCEL_FILE = "AI_Interview_Tracker.xlsx"

STATUS_MAP = {
    "applied": "Applied",
    "skipped_external": "Skipped - External",
    "already_applied_on_site": "Already Applied",
    "failed": "Failed",
}


def load_applied_ids() -> Set[str]:
    """Load previously applied job IDs from the CSV tracker."""
    ids: Set[str] = set()
    if not os.path.exists(CSV_FILE):
        return ids
    with open(CSV_FILE, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            ids.add(row["job_id"])
    return ids


def log_result(job: JobListing, status: str):
    """Append an application result to the CSV tracker."""
    file_exists = os.path.exists(CSV_FILE)
    with open(CSV_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        if not file_exists:
            writer.writeheader()
        writer.writerow(
            {
                "job_id": job.job_id,
                "title": job.title,
                "company": job.company,
                "link": job.link,
                "status": status,
                "timestamp": datetime.now().isoformat(),
            }
        )
    log_to_excel(job, status)


def log_external(job: JobListing):
    """Log an external-apply job to a separate CSV for manual follow-up."""
    file_exists = os.path.exists(EXTERNAL_CSV)
    with open(EXTERNAL_CSV, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=EXTERNAL_HEADERS)
        if not file_exists:
            writer.writeheader()
        writer.writerow({
            "title": job.title,
            "company": job.company,
            "link": job.link,
            "timestamp": datetime.now().isoformat(),
        })


def log_to_excel(job: JobListing, status: str):
    """Log application result to AI_Interview_Tracker.xlsx."""
    if not os.path.exists(EXCEL_FILE):
        return

    try:
        wb = load_workbook(EXCEL_FILE)
        today = datetime.now().strftime("%Y-%m-%d")
        display_status = STATUS_MAP.get(status, status)

        # Naukri Applications sheet
        apps_sheet = wb["Naukri Applications"]
        apps_sheet.append([
            today, job.company, job.title, job.link,
            "", "", "", "",  # Location, Experience, CTC, Skills
            display_status, "No", "", "", "",  # Status, Response, Date, Next, Notes
        ])

        # Dashboard — update Total Naukri Applications count (cell B3)
        dashboard = wb["Dashboard"]
        app_count = sum(
            1 for row in apps_sheet.iter_rows(min_row=2, max_col=1, values_only=True)
            if row[0]
        )
        dashboard["B3"] = app_count

        # Company Pipeline — add if company not already listed
        pipeline = wb["Company Pipeline"]
        existing_companies = set()
        for row in pipeline.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                existing_companies.add(row[0].strip().lower())

        if job.company.strip().lower() not in existing_companies:
            pipeline.append([
                job.company, job.title, today, "Applied",
                "", "", "", "",  # Next Interview, Prep, Contact, Notes
                "Active",
            ])

        wb.save(EXCEL_FILE)
    except Exception as exc:
        print(f"  [!] Excel logging failed: {exc}")


def _is_external_apply(page: Page) -> bool:
    """Check if the apply button leads to an external site."""
    external_selectors = [
        'a:has-text("Apply on company site")',
        'button:has-text("Apply on company site")',
        '[class*="company-site-btn"]',
        'a[target="_blank"]:has-text("Apply")',
    ]
    for sel in external_selectors:
        try:
            el = page.locator(sel).first
            if el.is_visible(timeout=1000):
                return True
        except (PlaywrightTimeout, Exception):
            continue
    return False


def _click_apply(page: Page) -> bool:
    """Find and click the direct apply button. Returns True if clicked."""
    apply_selectors = [
        'button:has-text("Apply")',
        'button[id*="apply"]',
        'button[class*="apply"]',
        '.apply-button',
        '#apply-button',
    ]
    for sel in apply_selectors:
        try:
            btn = page.locator(sel).first
            if btn.is_visible(timeout=1500):
                btn.click()
                return True
        except (PlaywrightTimeout, Exception):
            continue
    return False


def _handle_apply_modal(page: Page, config: dict) -> str:
    """Handle post-click modals like 'confirm resume' or chatbot questions.

    Returns: 'completed', 'partial', or 'no_modal'.
    """
    human_delay(0, 3.0)

    # Try chatbot first (delegated to chatbot.py)
    chatbot_result = handle_chatbot(page, config)
    if chatbot_result != "none":
        return chatbot_result

    submit_selectors = [
        'button:has-text("Submit")',
        'button:has-text("Apply")',
        'button:has-text("Confirm")',
        'button[type="submit"]',
    ]
    for sel in submit_selectors:
        try:
            btn = page.locator(sel).first
            if btn.is_visible(timeout=2000):
                btn.click()
                human_delay(0, 3.0)
                return "completed"
        except (PlaywrightTimeout, Exception):
            continue

    close_selectors = [
        'button:has-text("OK")',
        'button:has-text("Done")',
        ".crossIcon",
        '[class*="crossIcon"]',
    ]
    for sel in close_selectors:
        try:
            btn = page.locator(sel).first
            if btn.is_visible(timeout=1000):
                btn.click()
                human_delay(0, 3.0)
                return "completed"
        except (PlaywrightTimeout, Exception):
            continue

    return "no_modal"


def apply_to_job(page: Page, job: JobListing, config: dict) -> str:
    """Apply to a single job. Opens the job page, clicks apply, handles modals."""
    try:
        print(f"  [>] Opening: {job.title} at {job.company}")
        page.goto(job.link, wait_until="domcontentloaded")
        human_delay(0, 3.0)

        already_applied = page.locator(
            'text="Already Applied",'
            'text="You have already applied",'
            '[class*="already-applied"]'
        ).first
        try:
            if already_applied.is_visible(timeout=1500):
                print(f"  [=] Already applied on site")
                return "already_applied_on_site"
        except (PlaywrightTimeout, Exception):
            pass

        if _is_external_apply(page):
            print(f"  [~] External apply — saved to {EXTERNAL_CSV}")
            log_external(job)
            return "skipped_external"

        if _click_apply(page):
            modal_result = _handle_apply_modal(page, config)

            if modal_result == "partial":
                print(f"  [!] Chatbot incomplete — application NOT submitted")
                return "failed"

            # Verify application went through
            human_delay(0, 3.0)
            try:
                applied_check = page.locator(
                    'text="Already Applied",'
                    'text="You have already applied",'
                    '[class*="already-applied"],'
                    'text="Application Submitted",'
                    'text="applied successfully"'
                ).first
                if applied_check.is_visible(timeout=3000):
                    print(f"  [+] Applied successfully (verified)")
                    return "applied"
            except (PlaywrightTimeout, Exception):
                pass

            if modal_result == "completed":
                print(f"  [+] Applied (chatbot completed, no verify indicator)")
                return "applied"

            print(f"  [?] Apply clicked but couldn't verify — marking applied")
            return "applied"

        print(f"  [!] No apply button found")
        return "failed"

    except Exception as exc:
        print(f"  [!] Error applying: {exc}")
        return "failed"


def apply_to_jobs(
    page: Page,
    jobs: list[JobListing],
    config: dict,
    seen_jobs: set[tuple[str, str]] | None = None,
) -> dict:
    """Apply to a list of jobs. Returns a summary dict of counts."""
    applied_ids = load_applied_ids()
    if seen_jobs is None:
        seen_jobs = set()
    summary = {"applied": 0, "skipped_external": 0, "already_applied": 0, "failed": 0}
    skip_companies = [c.lower() for c in config.get("skip_companies", [])]
    required_keywords = [kw.lower() for kw in config.get("required_title_keywords", [])]

    for idx, job in enumerate(jobs, 1):
        print(f"\n[{idx}/{len(jobs)}] {job.title} — {job.company}")

        if job.job_id in applied_ids:
            print(f"  [=] Already in tracker — skipping")
            summary["already_applied"] += 1
            continue

        job_key = (job.title.strip().lower(), job.company.strip().lower())
        if job_key in seen_jobs:
            print(f"  [=] Duplicate listing — skipping")
            summary["already_applied"] += 1
            continue
        seen_jobs.add(job_key)

        if job.company.strip().lower() in skip_companies:
            print(f"  [x] Skipping blocked company: {job.company}")
            summary["failed"] += 1
            continue

        if required_keywords:
            title_lower = job.title.strip().lower()
            if not any(kw in title_lower for kw in required_keywords):
                print(f"  [x] Title doesn't match required keywords — skipping")
                summary["failed"] += 1
                continue

        status = apply_to_job(page, job, config)
        log_result(job, status)
        applied_ids.add(job.job_id)

        if status == "applied":
            summary["applied"] += 1
        elif status == "skipped_external":
            summary["skipped_external"] += 1
        elif status == "already_applied_on_site":
            summary["already_applied"] += 1
        else:
            summary["failed"] += 1

        human_delay(0, 3.0)

    return summary
